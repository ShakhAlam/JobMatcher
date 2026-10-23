import io
import zipfile

import pytest
from docx import Document
from fpdf import FPDF

from app.extract.documents import (
    MAX_TEXT_CHARS,
    MAX_UPLOAD_BYTES,
    ExtractionError,
    UnsupportedFileType,
    detect_type,
    extract_text,
    normalize_text,
)

RESUME_LINES = [
    "Alex Rivera - Senior Software Engineer",
    "8 years building Java and Spring Boot services for fintech.",
    "Skills: Java, Spring Boot, Hibernate, PostgreSQL, Kafka",
]


def make_pdf(lines=RESUME_LINES, blank_pages=0, user_password=None) -> bytes:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    if lines:
        pdf.add_page()
        for line in lines:
            pdf.cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
    for _ in range(blank_pages):
        pdf.add_page()
    if user_password:
        pdf.set_encryption(owner_password="owner", user_password=user_password)
    return bytes(pdf.output())


def make_docx(paragraphs=RESUME_LINES[:2], table_rows=(("Company", "PayStream"), ("Years", "2020-present"))) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=2)
        for row, values in zip(table.rows, table_rows):
            for cell, value in zip(row.cells, values):
                cell.text = value
    doc.add_paragraph("References available on request")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- unit -----------------------------------------------------------------


def test_normalize_text_collapses_whitespace():
    raw = "Name Here  \r\n\r\n\r\n\r\nSkills:\t\tJava   \nEnd\x00"
    assert normalize_text(raw) == "Name Here\n\nSkills: Java\nEnd"


def test_detects_type_from_content_not_name():
    assert detect_type(make_pdf(), "resume.docx") == "pdf"
    assert detect_type(make_docx(), "resume.pdf") == "docx"
    assert detect_type("plain résumé text".encode(), "resume") == "txt"


def test_rejects_legacy_doc():
    with pytest.raises(UnsupportedFileType, match="Legacy .doc"):
        detect_type(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100, "resume.doc")


def test_rejects_non_docx_zip_and_binary():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("hello.txt", "hi")
    with pytest.raises(UnsupportedFileType):
        detect_type(buf.getvalue(), "archive.docx")
    with pytest.raises(UnsupportedFileType):
        detect_type(b"\x89PNG\r\n\x1a\n\x00\x00", "photo.png")


def test_pdf_text_and_page_count():
    result = extract_text(make_pdf(), "resume.pdf")
    assert result.file_type == "pdf"
    assert result.pages == 1
    assert "Senior Software Engineer" in result.text
    assert "Kafka" in result.text
    assert result.warnings == []
    assert result.char_count == len(result.text)


def test_scanned_looking_pdf_warns():
    result = extract_text(make_pdf(lines=[], blank_pages=2), "scan.pdf")
    assert result.pages == 2
    assert result.text == ""
    assert "scanned" in result.warnings[0]


def test_password_protected_pdf_is_rejected():
    with pytest.raises(ExtractionError, match="password"):
        extract_text(make_pdf(user_password="secret"), "locked.pdf")


def test_corrupt_pdf_is_rejected():
    with pytest.raises(ExtractionError):
        extract_text(b"%PDF-1.7\nthis is not really a pdf", "broken.pdf")


def test_docx_keeps_paragraph_and_table_order():
    result = extract_text(make_docx(), "resume.docx")
    assert result.file_type == "docx"
    assert result.pages is None
    lines = result.text.splitlines()
    assert lines[0] == RESUME_LINES[0]
    assert "Company | PayStream" in lines
    assert lines.index("Company | PayStream") < lines.index("References available on request")


def test_empty_docx_warns():
    result = extract_text(make_docx(paragraphs=[], table_rows=()), "empty.docx")
    assert result.text == "References available on request"
    empty = Document()
    buf = io.BytesIO()
    empty.save(buf)
    assert "No text" in extract_text(buf.getvalue(), "blank.docx").warnings[0]


def test_docx_zip_bomb_is_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"0" * (51 * 1024 * 1024))
    with pytest.raises(ExtractionError, match="unreasonable size"):
        extract_text(buf.getvalue(), "bomb.docx")


def test_long_text_is_truncated_with_warning():
    result = extract_text(("word " * 20_000).encode(), "long.txt")
    assert result.char_count == MAX_TEXT_CHARS
    assert "cut" in result.warnings[0]


# --- API ------------------------------------------------------------------


def test_extract_endpoint_pdf(client):
    resp = client.post("/api/extract-text", files={"file": ("resume.pdf", make_pdf(), "application/pdf")})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["filename"] == "resume.pdf"
    assert body["file_type"] == "pdf"
    assert "Hibernate" in body["text"]


def test_extract_endpoint_docx(client):
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    resp = client.post("/api/extract-text", files={"file": ("cv.docx", make_docx(), mime)})
    assert resp.status_code == 200, resp.text
    assert "PayStream" in resp.json()["text"]


@pytest.mark.parametrize(
    ("payload", "status"),
    [
        (b"", 422),
        (b"\x89PNG\r\n\x1a\n\x00\x00", 415),
        (b"%PDF-1.7\ngarbage", 422),
        (b"a" * (MAX_UPLOAD_BYTES + 1), 413),
    ],
    ids=["empty", "png", "corrupt-pdf", "too-large"],
)
def test_extract_endpoint_errors(client, payload, status):
    resp = client.post("/api/extract-text", files={"file": ("file.bin", payload, "application/octet-stream")})
    assert resp.status_code == status, resp.text
    assert resp.json()["detail"]


def test_extract_endpoint_requires_file(client):
    assert client.post("/api/extract-text").status_code == 422
